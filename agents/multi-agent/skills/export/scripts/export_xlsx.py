from pathlib import Path
from openpyxl import Workbook, load_workbook


def export_rows(out_path: str, headers, rows):
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        wb = load_workbook(out)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    wb.save(out)
